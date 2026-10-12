#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
办公文件生成器（第二批计划 §3.2）。

============================================================================
 它是"受限结构化规格 → Office 文件"的**唯一**执行者
============================================================================
输入：一份 JSON 规格（**闭集**：见 `SPEC_SCHEMA` 与下面的校验）
输出：一个生成好的文件（写到 `--out` 指定的路径）+ 一份 JSON 读数（stdout 最后一行）

计划书 §3.2 逐字：「由模型提供**受限结构化内容**、可信模板和程序生成文件，
**不接受模型生成 Python/Shell 代码直接运行**」。
⇒ 本脚本**没有** `eval` / `exec` / `compile`，也不接受"要执行的代码"这类字段；
   它只认识下面固定几种块类型，别的键一律**拒收整条**。

三个库的固定版本由既有专属 venv 提供（python-docx 1.2.0 / python-pptx 1.0.2 /
openpyxl 3.1.5），本脚本不安装任何东西、不联网。

用法：
    python office-generate.py --format docx --spec <spec.json> --out <file>

stdout 的最后一行永远是**一个** JSON 对象：
    {"ok": true,  "bytes": 12345, "summary": {...}, "readback": {...}}
    {"ok": false, "errorType": "...", "error": "..."}
⚠️ 任何异常都被捕成 `ok:false`，绝不把 traceback 混进 stdout 的最后一行。
"""

import argparse
import json
import os
import sys

# ---------------------------------------------------------------------------
# 硬上界（与计划书 §3.2 的初始规模逐字一致）
# ---------------------------------------------------------------------------

LIMITS = {
    "docx_max_chars": 50_000,
    "xlsx_max_cells": 10_000,
    "pptx_max_slides": 20,
    "max_blocks": 500,
    "max_sheets": 10,
    "max_columns": 50,
    "max_rows": 500,
}

# 公式允许的函数名（**白名单**；不在这张表里的函数一律拒收）
ALLOWED_FORMULA_FUNCTIONS = {
    "SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "ROUND", "ABS", "IF",
}


class SpecError(Exception):
    """规格不合法（**用户/模型可修的错**，不是脚本的错）"""


def fail(error_type, message):
    print(json.dumps({"ok": False, "errorType": error_type, "error": message}, ensure_ascii=False))
    return 1


# ---------------------------------------------------------------------------
# 校验（**闭集**；未知字段整条拒收 —— 与 TS 侧 `parseTaskPlan` 同一条纪律）
# ---------------------------------------------------------------------------

def require_object(value, what):
    if not isinstance(value, dict):
        raise SpecError(f"{what} 必须是一个对象")
    return value


def reject_unknown(obj, allowed, what):
    for key in obj.keys():
        if key not in allowed:
            raise SpecError(f"{what} 有未知字段：{key}")


def require_text(value, what, limit=20_000):
    if not isinstance(value, str):
        raise SpecError(f"{what} 必须是字符串")
    if len(value) > limit:
        raise SpecError(f"{what} 长度 {len(value)} 超过上限 {limit}")
    return value


def check_formula(expr):
    """明确类型的公式：只允许白名单函数 + 单元格引用 + 数字 + 运算符。"""
    if not isinstance(expr, str) or not expr.startswith("="):
        raise SpecError("公式必须以 = 开头")
    body = expr[1:]
    allowed_chars = set("0123456789+-*/(),.:$ABCDEFGHIJKLMNOPQRSTUVWXYZ ")
    for ch in body.upper():
        if ch not in allowed_chars:
            raise SpecError(f"公式里有不允许的字符：{ch}")
    token = ""
    for ch in body.upper() + "(":
        if ch.isalpha() or ch == ".":
            token += ch
            continue
        if token:
            if ch == "(" and token not in ALLOWED_FORMULA_FUNCTIONS:
                raise SpecError(f"公式使用了白名单之外的函数：{token}")
            token = ""
    return expr


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

DOCX_BLOCKS = {"heading", "paragraph", "bullets", "table", "image"}
DOCX_TOP = {"title", "subtitle", "header", "footer", "blocks"}


def generate_docx(spec, out):
    from docx import Document
    from docx.shared import Pt

    reject_unknown(spec, DOCX_TOP, "docx 规格")
    blocks = spec.get("blocks", [])
    if not isinstance(blocks, list):
        raise SpecError("docx.blocks 必须是数组")
    if len(blocks) > LIMITS["max_blocks"]:
        raise SpecError(f"docx 块数 {len(blocks)} 超过上限 {LIMITS['max_blocks']}")

    total = 0
    document = Document()
    title = spec.get("title")
    if title is not None:
        text = require_text(title, "docx.title")
        total += len(text)
        document.add_heading(text, level=0)
    subtitle = spec.get("subtitle")
    if subtitle is not None:
        text = require_text(subtitle, "docx.subtitle")
        total += len(text)
        document.add_paragraph(text)

    headings = paragraphs = table_rows = images = bullets = 0
    for index, raw in enumerate(blocks):
        block = require_object(raw, f"docx.blocks[{index}]")
        kind = block.get("type")
        if kind == "heading":
            reject_unknown(block, {"type", "level", "text"}, f"docx.blocks[{index}]")
            level = block.get("level", 1)
            if not isinstance(level, int) or not (1 <= level <= 4):
                raise SpecError(f"docx.blocks[{index}].level 必须是 1..4 的整数")
            text = require_text(block.get("text"), f"docx.blocks[{index}].text")
            total += len(text)
            document.add_heading(text, level=level)
            headings += 1
        elif kind == "paragraph":
            reject_unknown(block, {"type", "text"}, f"docx.blocks[{index}]")
            text = require_text(block.get("text"), f"docx.blocks[{index}].text")
            total += len(text)
            document.add_paragraph(text)
            paragraphs += 1
        elif kind == "bullets":
            reject_unknown(block, {"type", "items"}, f"docx.blocks[{index}]")
            items = block.get("items")
            if not isinstance(items, list) or not items:
                raise SpecError(f"docx.blocks[{index}].items 必须是非空数组")
            for item in items:
                text = require_text(item, f"docx.blocks[{index}].items[]")
                total += len(text)
                document.add_paragraph(text, style="List Bullet")
                bullets += 1
        elif kind == "table":
            reject_unknown(block, {"type", "header", "rows"}, f"docx.blocks[{index}]")
            header = block.get("header")
            rows = block.get("rows")
            if not isinstance(header, list) or not header:
                raise SpecError(f"docx.blocks[{index}].header 必须是非空数组")
            if not isinstance(rows, list):
                raise SpecError(f"docx.blocks[{index}].rows 必须是数组")
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for col, cell in enumerate(header):
                value = require_text(cell, f"docx.blocks[{index}].header[]", 500)
                total += len(value)
                table.rows[0].cells[col].text = value
            for row in rows:
                if not isinstance(row, list):
                    raise SpecError(f"docx.blocks[{index}].rows[] 必须是数组")
                cells = table.add_row().cells
                for col in range(len(header)):
                    raw_value = row[col] if col < len(row) else ""
                    value = "" if raw_value is None else str(raw_value)
                    total += len(value)
                    cells[col].text = value
                table_rows += 1
        elif kind == "image":
            reject_unknown(block, {"type", "path", "widthInches"}, f"docx.blocks[{index}]")
            path = require_text(block.get("path"), f"docx.blocks[{index}].path", 1000)
            if not os.path.isfile(path):
                raise SpecError(f"docx.blocks[{index}].path 不是一个可读文件：{path}")
            width = block.get("widthInches")
            if width is not None and (not isinstance(width, (int, float)) or width <= 0 or width > 10):
                raise SpecError(f"docx.blocks[{index}].widthInches 必须是 0..10 的数字")
            if width is not None:
                document.add_picture(path, width=Pt(width * 72))
            else:
                document.add_picture(path)
            images += 1
        else:
            raise SpecError(f"docx.blocks[{index}].type 不支持：{kind}")

    if total > LIMITS["docx_max_chars"]:
        raise SpecError(f"docx 文本总量 {total} 字符超过上限 {LIMITS['docx_max_chars']}")

    header_text = spec.get("header")
    footer_text = spec.get("footer")
    if header_text is not None or footer_text is not None:
        section = document.sections[0]
        if header_text is not None:
            section.header.paragraphs[0].text = require_text(header_text, "docx.header", 500)
        if footer_text is not None:
            section.footer.paragraphs[0].text = require_text(footer_text, "docx.footer", 500)

    document.save(out)
    return {
        "kind": "docx",
        "headings": headings,
        "paragraphs": paragraphs,
        "bullets": bullets,
        "tableRows": table_rows,
        "images": images,
        "textChars": total,
    }, readback_docx(out, title, blocks)


def readback_docx(path, title, blocks):
    """**重新打开**生成的文件并按结构读回（"XML 可读"不等于"文件是好的"）。"""
    from docx import Document

    document = Document(path)
    paragraphs = [p.text for p in document.paragraphs if p.text != ""]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    images = 0
    try:
        images = len(document.inline_shapes)
    except Exception:
        images = -1
    # 逐个断言"规格里说要有的是否真的在文件里"
    wanted = []
    if title is not None and isinstance(title, str):
        wanted.append(title)
    for block in blocks:
        if not isinstance(block, dict):
            continue
        if block.get("type") in ("heading", "paragraph") and isinstance(block.get("text"), str):
            wanted.append(block["text"])
        if block.get("type") == "bullets" and isinstance(block.get("items"), list):
            wanted.extend([item for item in block["items"] if isinstance(item, str)])
    flat = "\n".join(paragraphs)
    for table in tables:
        for row in table:
            flat += "\n" + "\t".join(row)
    missing = [text for text in wanted if text not in flat]
    return {
        "paragraphCount": len(document.paragraphs),
        "tableCount": len(tables),
        "tableRows": sum(len(t) for t in tables),
        "inlineShapes": images,
        "wantedTexts": len(wanted),
        "missingTexts": missing[:5],
        "missingCount": len(missing),
    }


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

XLSX_TOP = {"sheets"}
XLSX_SHEET = {"name", "columns", "rows", "totals"}


def coerce_cell(value):
    """单元格值 → (写入值, 是否文本)。

    ⚠️ 计划书 §3.2 逐字：「普通文字以文字写入；公式采用**明确类型**及有限函数，
    **不把外部正文中的 `=` 自动当公式**」。
    ⇒ 只有 `{"f": "=SUM(...)"}` 这种**明确类型**才是公式；
       普通字符串以 `data_type='s'` 落盘，即使它以 `=` 开头。
    """
    if isinstance(value, dict):
        if set(value.keys()) != {"f"}:
            raise SpecError("单元格对象只允许有 f 一个字段（明确类型的公式）")
        return check_formula(value["f"]), "formula"
    if value is None:
        return None, "empty"
    if isinstance(value, bool):
        return value, "bool"
    if isinstance(value, (int, float)):
        return value, "number"
    return str(value), "text"


def generate_xlsx(spec, out):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    reject_unknown(spec, XLSX_TOP, "xlsx 规格")
    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise SpecError("xlsx.sheets 必须是非空数组")
    if len(sheets) > LIMITS["max_sheets"]:
        raise SpecError(f"xlsx 工作表数 {len(sheets)} 超过上限 {LIMITS['max_sheets']}")

    workbook = Workbook()
    workbook.remove(workbook.active)
    cells = 0
    totals_written = 0
    for index, raw in enumerate(sheets):
        sheet = require_object(raw, f"xlsx.sheets[{index}]")
        reject_unknown(sheet, XLSX_SHEET, f"xlsx.sheets[{index}]")
        name = require_text(sheet.get("name"), f"xlsx.sheets[{index}].name", 31)
        columns = sheet.get("columns")
        rows = sheet.get("rows", [])
        if not isinstance(columns, list) or not columns:
            raise SpecError(f"xlsx.sheets[{index}].columns 必须是非空数组")
        if len(columns) > LIMITS["max_columns"]:
            raise SpecError(f"xlsx.sheets[{index}] 列数超过上限 {LIMITS['max_columns']}")
        if not isinstance(rows, list) or len(rows) > LIMITS["max_rows"]:
            raise SpecError(f"xlsx.sheets[{index}].rows 必须是不超过 {LIMITS['max_rows']} 行的数组")

        worksheet = workbook.create_sheet(title=name)
        for col, header in enumerate(columns, start=1):
            cell = worksheet.cell(row=1, column=col)
            cell.value = require_text(header, f"xlsx.sheets[{index}].columns[]", 200)
            cell.font = Font(bold=True)
            cells += 1
        for row_index, row in enumerate(rows, start=2):
            if not isinstance(row, list):
                raise SpecError(f"xlsx.sheets[{index}].rows[] 必须是数组")
            for col in range(1, len(columns) + 1):
                value = row[col - 1] if col - 1 < len(row) else None
                written, kind = coerce_cell(value)
                cell = worksheet.cell(row=row_index, column=col)
                cell.value = written
                if kind == "text":
                    # ⚠️ 强制文本：openpyxl 会把以 = 开头的字符串当公式
                    cell.data_type = "s"
                cells += 1
        totals = sheet.get("totals", [])
        if not isinstance(totals, list):
            raise SpecError(f"xlsx.sheets[{index}].totals 必须是数组")
        for total in totals:
            total = require_object(total, f"xlsx.sheets[{index}].totals[]")
            reject_unknown(total, {"column", "label"}, f"xlsx.sheets[{index}].totals[]")
            column = total.get("column")
            if not isinstance(column, int) or not (1 <= column <= len(columns)):
                raise SpecError(f"xlsx.sheets[{index}].totals[].column 必须是 1..{len(columns)} 的整数")
            label = total.get("label", "合计")
            row_index = len(rows) + 2
            worksheet.cell(row=row_index, column=1).value = require_text(label, "totals[].label", 100)
            letter = worksheet.cell(row=1, column=column).column_letter
            formula = f"=SUM({letter}2:{letter}{len(rows) + 1})"
            cell = worksheet.cell(row=row_index, column=column)
            cell.value = formula
            cells += 1
            totals_written += 1
        if cells > LIMITS["xlsx_max_cells"]:
            raise SpecError(f"xlsx 单元格总数 {cells} 超过上限 {LIMITS['xlsx_max_cells']}")

    workbook.save(out)
    return {
        "kind": "xlsx",
        "sheets": len(sheets),
        "cells": cells,
        "totals": totals_written,
        "sheetNames": [s.get("name") for s in sheets if isinstance(s, dict)],
    }, readback_xlsx(out, sheets)


def readback_xlsx(path, sheets):
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False)
    names = list(workbook.sheetnames)
    formulas = 0
    texts_forced = 0
    wanted = []
    for sheet in sheets:
        if not isinstance(sheet, dict):
            continue
        for header in sheet.get("columns", []) or []:
            if isinstance(header, str):
                wanted.append(header)
    for name in names:
        worksheet = workbook[name]
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas += 1
                elif cell.data_type == "s":
                    texts_forced += 1
    flat = []
    for name in names:
        for row in workbook[name].iter_rows(values_only=True):
            flat.extend([str(v) for v in row if v is not None])
    joined = "\n".join(flat)
    missing = [text for text in wanted if text not in joined]
    return {
        "sheetNames": names,
        "formulaCells": formulas,
        "textCells": texts_forced,
        "wantedColumns": len(wanted),
        "missingColumns": missing[:5],
        "missingCount": len(missing),
    }


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

PPTX_TOP = {"title", "subtitle", "slides"}
PPTX_SLIDE = {"title", "bullets", "table", "chart"}


def generate_pptx(spec, out):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    reject_unknown(spec, PPTX_TOP, "pptx 规格")
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise SpecError("pptx.slides 必须是非空数组")
    if len(slides) > LIMITS["pptx_max_slides"]:
        raise SpecError(f"pptx 幻灯片数 {len(slides)} 超过上限 {LIMITS['pptx_max_slides']}")

    presentation = Presentation()
    title = spec.get("title")
    if title is not None:
        cover = presentation.slides.add_slide(presentation.slide_layouts[0])
        cover.shapes.title.text = require_text(title, "pptx.title", 200)
        subtitle = spec.get("subtitle")
        if subtitle is not None and len(cover.placeholders) > 1:
            cover.placeholders[1].text = require_text(subtitle, "pptx.subtitle", 400)

    bullets_total = 0
    tables = 0
    charts = 0
    for index, raw in enumerate(slides):
        slide_spec = require_object(raw, f"pptx.slides[{index}]")
        reject_unknown(slide_spec, PPTX_SLIDE, f"pptx.slides[{index}]")
        slide_title = require_text(slide_spec.get("title"), f"pptx.slides[{index}].title", 200)
        bullets = slide_spec.get("bullets")
        table = slide_spec.get("table")
        chart = slide_spec.get("chart")
        given = [x for x in (bullets, table, chart) if x is not None]
        if len(given) != 1:
            raise SpecError(f"pptx.slides[{index}] 必须有且只有 bullets / table / chart 之一")

        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = slide_title
        if bullets is not None:
            if not isinstance(bullets, list) or not bullets:
                raise SpecError(f"pptx.slides[{index}].bullets 必须是非空数组")
            body = slide.placeholders[1].text_frame
            body.clear()
            for position, item in enumerate(bullets):
                text = require_text(item, f"pptx.slides[{index}].bullets[]", 500)
                paragraph = body.paragraphs[0] if position == 0 else body.add_paragraph()
                paragraph.text = text
                bullets_total += 1
        elif table is not None:
            table = require_object(table, f"pptx.slides[{index}].table")
            reject_unknown(table, {"header", "rows"}, f"pptx.slides[{index}].table")
            header = table.get("header")
            rows = table.get("rows", [])
            if not isinstance(header, list) or not header:
                raise SpecError(f"pptx.slides[{index}].table.header 必须是非空数组")
            if not isinstance(rows, list):
                raise SpecError(f"pptx.slides[{index}].table.rows 必须是数组")
            shape = slide.shapes.add_table(
                rows=len(rows) + 1, cols=len(header), left=Inches(0.6), top=Inches(1.6),
                width=Inches(8.5), height=Inches(0.4 * (len(rows) + 1)),
            )
            grid = shape.table
            for col, value in enumerate(header):
                grid.cell(0, col).text = require_text(value, "table.header[]", 200)
            for row_index, row in enumerate(rows, start=1):
                if not isinstance(row, list):
                    raise SpecError("pptx table.rows[] 必须是数组")
                for col in range(len(header)):
                    raw_value = row[col] if col < len(row) else ""
                    grid.cell(row_index, col).text = "" if raw_value is None else str(raw_value)
            tables += 1
        else:
            chart = require_object(chart, f"pptx.slides[{index}].chart")
            reject_unknown(chart, {"kind", "categories", "series"}, f"pptx.slides[{index}].chart")
            kind = chart.get("kind")
            kinds = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE}
            if kind not in kinds:
                raise SpecError(f"pptx.slides[{index}].chart.kind 只支持 bar / line / pie")
            categories = chart.get("categories")
            series = chart.get("series")
            if not isinstance(categories, list) or not categories:
                raise SpecError(f"pptx.slides[{index}].chart.categories 必须是非空数组")
            if not isinstance(series, list) or not series:
                raise SpecError(f"pptx.slides[{index}].chart.series 必须是非空数组")
            data = CategoryChartData()
            data.categories = [str(c) for c in categories]
            for item in series:
                item = require_object(item, "chart.series[]")
                reject_unknown(item, {"name", "values"}, "chart.series[]")
                values = item.get("values")
                if not isinstance(values, list) or len(values) != len(categories):
                    raise SpecError("chart.series[].values 的长度必须与 categories 相同")
                for value in values:
                    if not isinstance(value, (int, float)) or isinstance(value, bool):
                        raise SpecError("chart.series[].values 只允许数字")
                data.add_series(require_text(item.get("name"), "chart.series[].name", 100),
                                tuple(float(v) for v in values))
            slide.shapes.add_chart(kinds[kind], Inches(0.6), Inches(1.6), Inches(8.5), Inches(4.5), data)
            charts += 1

    presentation.save(out)
    return {
        "kind": "pptx",
        "slides": len(slides) + (1 if title is not None else 0),
        "bullets": bullets_total,
        "tables": tables,
        "charts": charts,
    }, readback_pptx(out)


def readback_pptx(path):
    from pptx import Presentation

    presentation = Presentation(path)
    texts = []
    tables = 0
    charts = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs)
                    if line:
                        texts.append(line)
            if getattr(shape, "has_table", False):
                tables += 1
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text:
                            texts.append(cell.text)
            if getattr(shape, "has_chart", False):
                charts += 1
    return {
        "slideCount": len(presentation.slides),
        "textShapes": len(texts),
        "tables": tables,
        "charts": charts,
        "textSample": texts[:3],
    }


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------

GENERATORS = {
    "docx": generate_docx,
    "xlsx": generate_xlsx,
    "pptx": generate_pptx,
}


def main(argv):
    parser = argparse.ArgumentParser(description="受限结构化规格 → Office 文件")
    parser.add_argument("--format", required=True, choices=sorted(GENERATORS.keys()))
    parser.add_argument("--spec", required=True, help="JSON 规格文件路径")
    parser.add_argument("--out", required=True, help="输出文件路径（由调用方给临时路径）")
    args = parser.parse_args(argv)

    try:
        with open(args.spec, "r", encoding="utf-8") as handle:
            spec = json.load(handle)
    except Exception as error:  # noqa: BLE001
        return fail("bad_spec_file", f"读不到 JSON 规格：{error}")

    if not isinstance(spec, dict):
        return fail("bad_spec", "规格必须是一个 JSON 对象")

    try:
        summary, readback = GENERATORS[args.format](spec, args.out)
    except SpecError as error:
        return fail("spec_rejected", str(error))
    except Exception as error:  # noqa: BLE001
        return fail("generate_failed", f"{type(error).__name__}: {error}")

    size = os.path.getsize(args.out) if os.path.isfile(args.out) else -1
    print(json.dumps({"ok": True, "bytes": size, "summary": summary, "readback": readback},
                     ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
